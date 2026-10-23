"""Build a text sales report from an XLSX ranking sheet and its cell notes."""

import argparse
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import sys
import tempfile
from zipfile import BadZipFile

import openpyxl


@dataclass(frozen=True)
class Agent:
    name: str
    sales: int
    notes: str = ""


def confirmed_count(notes):
    """Read complete '12 confirmed' or 'confirmed' segments; ignore other labels."""
    total = 0
    for segment in re.split(r"[;,\n|]", notes or ""):
        segment = segment.strip()
        match = re.fullmatch(r"(?:(\d+)\s+)?confirmed", segment, re.IGNORECASE)
        if match:
            total += int(match.group(1)) if match.group(1) is not None else 1
        elif re.search(r"\bconfirmed\b", segment, re.IGNORECASE):
            raise ValueError(f"Unsupported note segment {segment!r}; use '12 confirmed' or 'confirmed'.")
    return total


def sales_count(value, location):
    if isinstance(value, bool) or value is None or str(value).strip() == "":
        raise ValueError(f"{location}: sales must be a nonnegative whole number (or '-' to skip).")
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        number = Decimal("NaN")
    if not number.is_finite() or number < 0 or number != number.to_integral_value():
        raise ValueError(f"{location}: invalid sales value {value!r}; expected a nonnegative whole number.")
    return int(number)


def campaign_columns(sheet):
    columns = {}
    for col in range(1, sheet.max_column):
        name = sheet.cell(1, col).value
        labels = [str(sheet.cell(2, c).value or "").strip().casefold() for c in (col, col + 1)]
        if isinstance(name, str) and name.strip() and labels == ["agent", "sales"]:
            key = name.strip().casefold()
            if key in columns:
                raise ValueError(f"Duplicate campaign heading {name.strip()!r} in sheet {sheet.title!r}.")
            columns[key] = (name.strip(), col)
    if not columns:
        raise ValueError(f"No campaign blocks in {sheet.title!r}. Expected row 1: campaign; row 2: Agent, Sales.")
    return columns


def read_agents(sheet, column):
    agents = []
    seen = set()
    for row in range(3, sheet.max_row + 1):
        name = str(sheet.cell(row, column).value or "").strip()
        if name.casefold() == "total":
            break
        if not name:
            continue
        cell = sheet.cell(row, column + 1)
        if isinstance(cell.value, str) and cell.value.strip() == "-":
            continue
        if name.casefold() in seen:
            raise ValueError(f"{sheet.title}!{cell.coordinate}: duplicate agent {name!r} in a campaign.")
        seen.add(name.casefold())
        notes = cell.comment.text.strip() if cell.comment else ""
        try:
            confirmed_count(notes)
        except ValueError as error:
            raise ValueError(f"{sheet.title}!{cell.coordinate}: {error}") from error
        agents.append(Agent(name, sales_count(cell.value, f"{sheet.title}!{cell.coordinate}"), notes))
    return agents


def generate_report(file_path, sheet_name=None, campaigns=None):
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    try:
        if sheet_name is None:
            if len(workbook.sheetnames) != 1:
                raise ValueError("Choose --sheet explicitly. Available sheets: " + ", ".join(workbook.sheetnames))
            sheet_name = workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet {sheet_name!r} not found. Available sheets: {', '.join(workbook.sheetnames)}")
        sheet = workbook[sheet_name]
        columns = campaign_columns(sheet)
        selected = list(columns) if campaigns is None else [name.strip().casefold() for name in campaigns]
        if not selected or len(selected) != len(set(selected)):
            raise ValueError("Select at least one campaign, without duplicate names.")
        sections = []
        total_sales = total_confirmed = 0
        for key in selected:
            if key not in columns:
                raise ValueError(f"Campaign {key!r} not found in {sheet_name!r}. Available: {', '.join(c[0] for c in columns.values())}")
            name, column = columns[key]
            agents = read_agents(sheet, column)
            sales = sum(agent.sales for agent in agents)
            confirmed = sum(confirmed_count(agent.notes) for agent in agents)
            lines = [f"{name} | Sales: {sales} | Confirmed: {confirmed}"]
            for agent in agents:
                notes = " ".join(agent.notes.split()) or "No notes"
                lines.append(f"  {agent.name}: {agent.sales} ({notes})")
            sections.append("\n".join(lines))
            total_sales += sales
            total_confirmed += confirmed
        return (f"SALES RANKING REPORT\nSheet: {sheet_name}\n"
                f"Total sales: {total_sales}\nTotal confirmed: {total_confirmed}\n\n"
                + "\n\n".join(sections) + "\n")
    finally:
        workbook.close()


def write_report(output, text):
    """Replace an existing report only after successful validation and writing."""
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", newline="\n",
                                         dir=output.parent, delete=False, suffix=".tmp") as handle:
            temporary = Path(handle.name)
            handle.write(text)
        temporary.replace(output)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Path to the XLSX ranking workbook")
    parser.add_argument("--sheet", help="Exact sheet name; required for workbooks with multiple sheets")
    parser.add_argument("--campaign", action="append", help="Campaign to include; repeat for several (default: all)")
    parser.add_argument("--output", type=Path, help="Text output path (default: ranking-report.txt beside the workbook)")
    args = parser.parse_args(argv)
    if args.workbook.suffix.casefold() != ".xlsx":
        parser.error("The input must be an .xlsx workbook.")
    output = args.output or args.workbook.with_name("ranking-report.txt")
    if output.suffix.casefold() != ".txt":
        parser.error("--output must end in .txt")
    try:
        report = generate_report(args.workbook, args.sheet, args.campaign)
        write_report(output, report)
    except (ValueError, OSError, BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1
    print(f"Report saved to {output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
